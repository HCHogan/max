#!/usr/bin/env python3
"""Download the four Destiny guide workbooks as one CSV per worksheet.

Requires openpyxl and beautifulsoup4. Uses public XLSX exports where available,
otherwise reads the publicly rendered HTML worksheets. Numbers are written as
the sheet displays them (percentages, rounding), without thousands separators.
Writes en/ and archive/en/ grids, meta/ sidecars and .state/sources.json (see
destiny_guide_files.py); run translate-destiny-guides.py afterwards. The
current Chinese copies are harvested into the translation cache first, so
unchanged text is not translated again.
"""

import argparse
import json
import math
import re
import shutil
import subprocess
import tempfile
import warnings
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timezone
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from urllib.parse import urljoin
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from destiny_guide_files import SOURCES_INDEX, digest, grid_bytes, harvest_memory, open_memory


ROOT = Path(__file__).resolve().parent.parent
SOURCES = [
    ("court", "Destiny 2 Damage Buffs/Debuffs - JUNE UPDATE", "1i1KUwgVkd8qhwYj481gkV9sZNJQCE-C3Q-dpQutPCi4"),
    ("compendium", "Destiny Data Compendium", "1WaxvbLx7UoSZaBqdFr1u32F2uWVLo-CJunJB4nlGUE4"),
    ("endgame-analysis", "Destiny 2: Endgame Analysis", "1JM-0SlxVDAi-C6rGVlLxa-J1WGewEeL8Qvq4htWZHhY"),
    ("boss-damage", "Destiny 2: Boss Damage", "1_5wtBjRYHHxuF4oJKDb_iOGZs-wTkzB6RYbnyNLbuz4"),
]
# Generated paths this script owns; README.md and the translation outputs stay.
MANAGED = ["en", "archive/en", "meta", SOURCES_INDEX.as_posix()]
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
FIXED_FORMAT = re.compile(r"[#0,]*(?:\.([0#]+))?(%?)")


def stamp():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def download(url, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    result = subprocess.run(
        ["curl", "--location", "--fail", "--silent", "--show-error", "--max-time", "90",
         "--retry", "2", "--retry-delay", "1", url, "--output", str(path)],
        capture_output=True, text=True, check=False,
    )
    if result.returncode:
        raise RuntimeError(result.stderr.strip())
    return path.read_bytes()


def discover(html):
    """Extract the public tab inventory, including hidden tabs when exposed."""
    match = re.search(r"bootstrapData\s*=\s*(\{.*?\});\s*function loadWaffle", html, re.S)
    if match:
        snapshot = json.loads(match[1])
        tabs = []
        for operation, payload in snapshot["changes"]["topsnapshot"]:
            if operation != 21350203:
                continue
            values = json.loads(payload)
            properties = values[3][0]["1"]
            hidden = any(isinstance(p, dict) and p.get("9") == 1 for p in properties)
            tabs.append({"gid": str(values[2]), "title": properties[0][2],
                         "visibility": "hidden" if hidden else "visible"})
        if tabs:
            return tabs
    tabs = [{"title": name, "gid": gid, "visibility": "unknown"}
            for name, gid in re.findall(r'items\.push\(\{name: "([^"]*)"[^}]*gid: "(\d+)"', html)]
    if not tabs:
        raise ValueError("No worksheet inventory in the public document")
    return tabs


def general_number(value):
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
        return str(int(value))
    return format(Decimal(repr(value)), "f")


def number_text(value, number_format):
    """The sheet's displayed number: fixed decimals and percent, no grouping.

    Every format in the four workbooks is a single fixed-point section such as
    0.000, #,##, 0.00% or General; anything else keeps the full value.
    """
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    match = FIXED_FORMAT.fullmatch(number_format or "")
    if number_format == "General" or match is None or not re.search(r"[0#]", number_format):
        return general_number(value)
    places, percent = match[1] or "", bool(match[2])
    scaled = Decimal(repr(value)) * (100 if percent else 1)
    text = format(scaled.quantize(Decimal(1).scaleb(-len(places)), ROUND_HALF_UP), "f")
    optional = len(places) - len(places.rstrip("#"))
    if optional and "." in text:
        text = text[: len(text) - optional] + text[len(text) - optional:].rstrip("0")
        text = text.rstrip(".")
    if re.fullmatch(r"-0(\.0*)?", text):
        text = text[1:]
    return text + ("%" if percent else "")


def cell_text(value, number_format):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time().isoformat() == "00:00:00" else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        return number_text(value, number_format)
    return str(value)


def html_grid(raw, gid, source_url):
    soup = BeautifulSoup(raw, "html.parser")
    table = soup.find("table", class_="waffle")
    if table is None:
        raise ValueError(f"No worksheet table: {source_url}")
    rendered = []
    for tr in table.select("tbody > tr"):
        header = tr.find("th", recursive=False)
        match = re.fullmatch(re.escape(gid) + r"R(\d+)", header.get("id", "")) if header else None
        if match:
            rendered.append((int(match[1]) + 1, tr))
    if not rendered:
        raise ValueError(f"No numbered worksheet rows: {source_url}")
    grid, occupied = {}, set()
    metadata = {"merged_ranges": [], "hyperlinks": {}, "cell_notes": {},
                "rendered_rows": [row for row, _ in rendered], "has_images": False}
    columns = len(table.select("thead th[id]"))
    for position, (row, tr) in enumerate(rendered):
        col = 1
        for td in tr.find_all("td", recursive=False):
            while (row, col) in occupied:
                col += 1
            colspan, rowspan = int(td.get("colspan", 1)), int(td.get("rowspan", 1))
            covered = rendered[position:position + rowspan]
            for next_row, _ in covered:
                for next_col in range(col, col + colspan):
                    occupied.add((next_row, next_col))
            ref = f"{get_column_letter(col)}{row}"
            if colspan > 1 or rowspan > 1:
                end_row = covered[-1][0]
                metadata["merged_ranges"].append(f"{ref}:{get_column_letter(col + colspan - 1)}{end_row}")
            links = [{"text": a.get_text(), "url": urljoin(source_url, a["href"])}
                     for a in td.find_all("a", href=True)]
            if links:
                metadata["hyperlinks"][ref] = links
            if td.find("img", src=True):
                metadata["has_images"] = True
            for br in td.find_all("br"):
                br.replace_with("\n")
            grid[row, col] = td.get_text()
            columns = max(columns, col + colspan - 1)
            col += colspan
    rows = [[grid.get((r, c), "") for c in range(1, columns + 1)]
            for r in range(1, rendered[-1][0] + 1)]
    metadata["notes_availability"] = "Public HTML does not expose all spreadsheet cell notes or formulas."
    return rows, metadata


def xlsx_metadata(archive, number):
    path = f"xl/worksheets/sheet{number}.xml"
    root = ET.fromstring(archive.read(path))
    metadata = {"merged_ranges": [node.attrib["ref"] for node in root.findall("s:mergeCells/s:mergeCell", NS)],
                "hyperlinks": {}, "cell_notes": {}, "formulas": {},
                "hidden_rows": [int(node.attrib["r"]) for node in root.findall("s:sheetData/s:row", NS)
                                if node.get("hidden") == "1"]}
    relpath = f"xl/worksheets/_rels/sheet{number}.xml.rels"
    rels = {}
    if relpath in archive.namelist():
        for node in ET.fromstring(archive.read(relpath)):
            rels[node.attrib["Id"]] = node.attrib["Target"]
    for node in root.findall("s:hyperlinks/s:hyperlink", NS):
        rid = node.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        metadata["hyperlinks"][node.attrib["ref"]] = rels.get(rid, node.get("location", ""))
    for cell in root.findall("s:sheetData/s:row/s:c", NS):
        formula = cell.find("s:f", NS)
        if formula is not None:
            metadata["formulas"][cell.attrib["r"]] = formula.text or ""
    for target in rels.values():
        if "comments" not in target.lower() or not target.endswith(".xml"):
            continue
        import posixpath
        target = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl/worksheets", target))
        if target not in archive.namelist():
            continue
        for comment in ET.fromstring(archive.read(target)).findall("s:commentList/s:comment", NS):
            metadata["cell_notes"][comment.attrib["ref"]] = "".join(comment.itertext())
    metadata["has_images"] = root.find("s:drawing", NS) is not None
    return metadata


def sheet_status(tab):
    title = tab["title"].lower()
    return "ignore" if "ignore" in title else "historical" if re.search(r"\bold\b|outdated|not updated|historical", title) else "unspecified"


def write_sheet(stage, source_id, tab, slug, rows, metadata, retrieval):
    # Trim only trailing empty rows; line N of the CSV remains spreadsheet row N.
    while rows and not any(rows[-1]):
        rows.pop()
    columns = max((len(row) for row in rows), default=0)
    rows = [row + [""] * (columns - len(row)) for row in rows]
    status = sheet_status(tab)
    current = tab["visibility"] == "visible" and status == "unspecified"
    relative = Path("en" if current else "archive/en") / source_id / f"{slug}.csv"
    data = grid_bytes(rows)
    (stage / relative).parent.mkdir(parents=True, exist_ok=True)
    (stage / relative).write_bytes(data)
    meta = Path("meta") / source_id / f"{slug}.json"
    (stage / meta).parent.mkdir(parents=True, exist_ok=True)
    (stage / meta).write_text(json.dumps(metadata, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    vocabulary = {"name", "type", "family", "notes", "description", "perk", "perk 1", "perk 2", "icon",
                  "armor", "effect", "trigger", "tier", "source", "frame", "bonus", "pcs", "date",
                  "tab", "updated", "status", "activity", "boss", "total", "dps", "energy", "season"}
    headers = [number for number, row in enumerate(rows, 1)
               if sum(value.strip().lower() in vocabulary for value in row) >= 2][:3]
    return {"gid": tab["gid"], "title": tab["title"], "visibility": tab["visibility"], "source_status": status,
            "current": current, "path": relative.as_posix(), "meta_path": meta.as_posix(),
            "source_url": retrieval["source_url"], "retrieval_method": retrieval["method"],
            "fetched_at": retrieval["fetched_at"], "rows": len(rows), "columns": columns,
            "nonempty_cells": sum(bool(value) for row in rows for value in row),
            "header_rows": headers, "sha256": digest(data)}


def slugs(tabs):
    """Title slugs, with the gid appended only where two titles collide."""
    base = {tab["gid"]: re.sub(r"[^a-z0-9]+", "-", tab["title"].lower()).strip("-") or "sheet" for tab in tabs}
    counts = {}
    for value in base.values():
        counts[value] = counts.get(value, 0) + 1
    return {gid: value if counts[value] == 1 else f"{value}-{gid}" for gid, value in base.items()}


def fetch_source(source, cache, stage):
    source_id, title, spreadsheet_id = source
    base = f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}"
    html = download(base + "/edit", cache / source_id / "catalog.html").decode("utf-8")
    tabs = discover(html)
    names = slugs(tabs)
    workbook_path = cache / source_id / "workbook.xlsx"
    export_error = None
    try:
        download(base + "/export?format=xlsx", workbook_path)
        method = "xlsx_display_values"
    except RuntimeError as error:
        export_error = str(error)
        method = "public_html"
    source_stamp = stamp()
    records = []
    if method == "xlsx_display_values":
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            workbook = load_workbook(workbook_path, read_only=True, data_only=True)
        by_name = {tab["title"]: tab for tab in tabs}
        if set(by_name) != set(workbook.sheetnames):
            raise ValueError(f"Worksheet inventory differs from XLSX: {source_id}")
        with ZipFile(workbook_path) as archive:
            for number, sheet in enumerate(workbook, 1):
                tab = {**by_name[sheet.title], "visibility": sheet.sheet_state}
                rows = [[cell_text(cell.value, getattr(cell, "number_format", "General")) for cell in row]
                        for row in sheet.iter_rows()]
                metadata = xlsx_metadata(archive, number)
                records.append(write_sheet(stage, source_id, tab, names[tab["gid"]], rows, metadata,
                    {"source_url": base + f"/edit?gid={tab['gid']}", "method": method, "fetched_at": source_stamp}))
        workbook.close()
    else:
        def fetch_tab(tab):
            url = base + f"/htmlview/sheet?headers=true&gid={tab['gid']}"
            raw = download(url, cache / source_id / f"{tab['gid']}.html")
            rows, metadata = html_grid(raw, tab["gid"], base + f"/edit?gid={tab['gid']}")
            return write_sheet(stage, source_id, tab, names[tab["gid"]], rows, metadata,
                {"source_url": base + f"/edit?gid={tab['gid']}", "method": method, "fetched_at": stamp()})
        with ThreadPoolExecutor(max_workers=4) as pool:
            records = list(pool.map(fetch_tab, tabs))
    print(f"{source_id}: {len(records)} CSVs, {sum(record['rows'] for record in records)} rows ({method})", flush=True)
    return {"id": source_id, "title": title, "spreadsheet_id": spreadsheet_id, "source_url": base + "/edit",
            "fetched_at": source_stamp, "retrieval_method": method, "direct_export_error": export_error, "sheets": records}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=ROOT / "data/destiny-guides")
    parser.add_argument("--cache", type=Path, default=ROOT / "var/destiny-guides-import")
    parser.add_argument("--translation-cache", type=Path, default=ROOT / "var/destiny-guides-translation")
    args = parser.parse_args()
    with tempfile.TemporaryDirectory(prefix="max-destiny-guides-") as directory:
        stage = Path(directory)
        # Sources are sequential; only independent public worksheet downloads run in parallel.
        sources = [fetch_source(source, args.cache, stage) for source in SOURCES]
        index = {"format_version": 2, "generated_at": stamp(), "sources": sources}
        (stage / SOURCES_INDEX).parent.mkdir(parents=True, exist_ok=True)
        (stage / SOURCES_INDEX).write_text(json.dumps(index, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
        # Complete all downloads and parsing before replacing any checked-in data.
        db = open_memory(args.translation_cache)
        print(f"Harvested {harvest_memory(args.output, db)} translated lines into {args.translation_cache}")
        db.close()
        for relative in MANAGED:
            target = args.output / relative
            if target.is_dir():
                shutil.rmtree(target)
            elif target.exists():
                target.unlink()
        shutil.copytree(stage, args.output, dirs_exist_ok=True)
        print(f"Wrote {sum(len(source['sheets']) for source in sources)} CSVs to {args.output}; run translate-destiny-guides.py next")


if __name__ == "__main__":
    main()
